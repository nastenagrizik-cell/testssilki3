from __future__ import annotations

import re
from dataclasses import dataclass, field
from pathlib import Path

from docx import Document
from docx.document import Document as DocumentObject
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph
from docx.text.run import Run

from .logic_parser import looks_like_instruction, parse_logic_rules, parse_rotation
from .models import (
    AnswerOption,
    QuestionSpec,
    QuestionType,
    SourceLocation,
    SurveySpec,
)
from .normalization import normalize_id, normalize_space


QUESTION_ID_TOKEN = r"[A-Za-zА-Яа-яЁё]{1,5}\d+(?:[._]\d+|[A-Za-zА-Яа-яЁё]+)*"
QUESTION_PATTERN = re.compile(
    rf"^\s*(?P<id>{QUESTION_ID_TOKEN})\s*[.)]?\s*[:.-]?\s*(?P<text>.+)$"
)
ANSWER_CODE_PATTERN = re.compile(r"^\s*(\d{1,3}(?:\.\d+)?)\s*$")


@dataclass
class Block:
    index: int
    kind: str
    text: str
    raw_text: str = ""
    style: str = ""
    flags: set[str] = field(default_factory=set)
    cells: list[str] = field(default_factory=list)
    raw_cells: list[str] = field(default_factory=list)
    table_id: int | None = None


def _iter_body(doc: DocumentObject):
    for child in doc.element.body.iterchildren():
        if child.tag == qn("w:p"):
            yield Paragraph(child, doc)
        elif child.tag == qn("w:tbl"):
            yield Table(child, doc)


def _paragraph_flags(paragraph: Paragraph) -> set[str]:
    flags: set[str] = set()
    for run in paragraph.runs:
        if not normalize_space(run.text):
            continue
        if _run_is_struck(run):
            continue
        if run.bold:
            flags.add("bold")
        if run.italic:
            flags.add("italic")
        if run.underline:
            flags.add("underline")
        if run.font.color.rgb:
            flags.add(f"color:{run.font.color.rgb}")
        if run.font.highlight_color:
            flags.add(f"highlight:{run.font.highlight_color}")
    return flags


def _xml_switch_is_on(element) -> bool:
    if element is None:
        return False
    value = element.get(qn("w:val"))
    return value is None or value.lower() not in {"0", "false", "off", "none"}


def _run_is_struck(run: Run) -> bool:
    """Return True for direct or style-based single/double strikethrough."""
    if run.font.strike is True or run.font.double_strike is True:
        return True
    try:
        if run.style and (run.style.font.strike is True or run.style.font.double_strike is True):
            return True
    except (AttributeError, KeyError):
        pass
    properties = run._r.rPr
    return bool(
        properties is not None
        and (
            _xml_switch_is_on(properties.find(qn("w:strike")))
            or _xml_switch_is_on(properties.find(qn("w:dstrike")))
        )
    )


def _inside_deleted_change(element) -> bool:
    parent = element.getparent()
    while parent is not None:
        if parent.tag == qn("w:del"):
            return True
        parent = parent.getparent()
    return False


def _paragraph_text(paragraph: Paragraph, include_struck: bool = False) -> str:
    """Read visible Word text, omitting tracked deletions and struck runs."""
    parts: list[str] = []
    for run_element in paragraph._p.iter(qn("w:r")):
        if _inside_deleted_change(run_element):
            continue
        run = Run(run_element, paragraph)
        if not include_struck and _run_is_struck(run):
            continue
        parts.append(run.text)
    return normalize_space("".join(parts))


def extract_blocks(path: str | Path) -> list[Block]:
    doc = Document(path)
    blocks: list[Block] = []
    index = 0
    table_id = 0
    for item in _iter_body(doc):
        if isinstance(item, Paragraph):
            raw_text = _paragraph_text(item, include_struck=True)
            text = _paragraph_text(item)
            if raw_text or text:
                blocks.append(
                    Block(
                        index=index,
                        kind="paragraph",
                        text=text,
                        raw_text=raw_text,
                        style=item.style.name if item.style else "",
                        flags=_paragraph_flags(item),
                    )
                )
                index += 1
        else:
            table_id += 1
            for row in item.rows:
                cells = [normalize_space(" ".join(_paragraph_text(p) for p in cell.paragraphs)) for cell in row.cells]
                raw_cells = [
                    normalize_space(" ".join(_paragraph_text(p, include_struck=True) for p in cell.paragraphs))
                    for cell in row.cells
                ]
                if any(raw_cells) or any(cells):
                    blocks.append(
                        Block(
                            index=index,
                            kind="table_row",
                            text=" | ".join(cell for cell in cells if cell),
                            raw_text=" | ".join(cell for cell in raw_cells if cell),
                            cells=cells,
                            raw_cells=raw_cells,
                            table_id=table_id,
                        )
                    )
                    index += 1
    return blocks


def _parse_question_candidate(candidate: str) -> tuple[str, str] | None:
    candidate = normalize_space(candidate)
    if not candidate:
        return None
    match = QUESTION_PATTERN.match(candidate)
    if not match:
        return None
    qid = normalize_id(match.group("id"))
    text = normalize_space(match.group("text"))
    if not any(char.isalpha() for char in text):
        return None
    if re.fullmatch(
        r"(?:[A-Za-zА-Яа-яЁё]\d+(?:[._]\d+)?\s*)?(?:[<>=]+\s*[A-Za-zА-Яа-яЁё]\d+(?:[._]\d+)?\s*)+",
        text,
    ):
        return None
    if text.upper().startswith(("КОД ", "РОТАЦ", "ВАРИАНТ ")):
        return None
    return qid, text


def _question_from_block(block: Block) -> tuple[str, str] | None:
    active_candidates = block.cells if block.cells else [block.text]
    raw_candidates = block.raw_cells if block.raw_cells else [block.raw_text or block.text]
    for raw_candidate, active_candidate in zip(raw_candidates, active_candidates):
        if block.kind == "table_row" and len(
            re.findall(rf"(?:^|\s)({QUESTION_ID_TOKEN})\s*:", raw_candidate)
        ) > 1:
            continue
        raw_parsed = _parse_question_candidate(raw_candidate)
        if not raw_parsed or not normalize_space(active_candidate):
            continue
        active_parsed = _parse_question_candidate(active_candidate)
        if active_parsed:
            return active_parsed
        qid = raw_parsed[0]
        # The code and punctuation may live in a separate run.  Remove them
        # from the active wording when only a fragment of the question was struck.
        text = re.sub(
            rf"^\s*{re.escape(qid)}\s*[.)]?\s*[:.-]?\s*",
            "",
            normalize_space(active_candidate),
            flags=re.I,
        )
        if not text or not any(char.isalpha() for char in text):
            continue
        # A valid question can itself contain words such as "если",
        # "продолжите" or "ротация".  Those words are instructions only when
        # the block does not begin with a question code.
        return qid, text
    return None


def _is_question_boundary(block: Block) -> bool:
    candidates = block.raw_cells if block.raw_cells else [block.raw_text or block.text]
    for candidate in candidates:
        if block.kind == "table_row" and len(
            re.findall(rf"(?:^|\s)({QUESTION_ID_TOKEN})\s*:", candidate)
        ) > 1:
            continue
        if _parse_question_candidate(candidate):
            return True
    return False


def _infer_type(context: str, has_options: bool) -> QuestionType:
    upper = context.upper().replace("Ё", "Е")
    if "ИНФОБЛОК" in upper or "ВВОДНЫЙ ТЕКСТ" in upper:
        return QuestionType.INFO
    if "РАНЖ" in upper:
        return QuestionType.RANKING
    if "ТОЛЬКО ЦИФР" in upper or "ПРИ КАКОЙ ЦЕНЕ" in upper:
        return QuestionType.NUMBER
    if re.search(r"ОТКРЫТ(?:ЫЙ|ОГО|АЯ)?\s+ОТВЕТ", upper) or "ПОЛЕ ДЛЯ ВВОДА" in upper or "ЗАПИШИТЕ" in upper:
        return QuestionType.TEXT
    if "МАТРИЦ" in upper and ("НЕСКОЛЬК" in upper or "МНОЖЕСТВ" in upper):
        return QuestionType.MATRIX_MULTI
    if "МАТРИЦ" in upper or "ПО КАЖДОЙ СТРОК" in upper:
        return QuestionType.MATRIX_SINGLE
    if "ОДИН ОТВЕТ" in upper or "ЕДИНИЧН" in upper:
        return QuestionType.SINGLE
    if "НЕСКОЛЬК" in upper or "МНОЖЕСТВ" in upper:
        return QuestionType.MULTI
    if "ШКАЛ" in upper and has_options:
        return QuestionType.SCALE
    return QuestionType.SINGLE if has_options else QuestionType.UNKNOWN


def _make_answer(code: str, text: str) -> AnswerOption:
    upper = text.upper().replace("Ё", "Е")
    return AnswerOption(
        code=code,
        text=text,
        exclusive=(
            "ИСКЛЮЧА" in upper
            or "НЕ МОЖЕТ БЫТЬ ВЫБРАН" in upper
            or upper.startswith("НИЧЕГО ИЗ ПЕРЕЧИСЛЕННОГО")
            or upper.startswith("НЕТ ТАКИХ")
        ),
        fixed="ЗАКРЕПЛ" in upper,
    )


def _answers_from_block(block: Block) -> list[AnswerOption]:
    if not block.cells:
        return []
    answers: list[AnswerOption] = []
    if len(block.cells) >= 2:
        code = normalize_space(block.cells[0])
        text = normalize_space(block.cells[1])
        if ANSWER_CODE_PATTERN.match(code) and text:
            answers.append(_make_answer(code, text))
            return answers
    for cell in block.cells:
        match = re.match(r"^\s*(\d{1,3}(?:\.\d+)?)\s+(.+)$", normalize_space(cell))
        if match:
            answers.append(_make_answer(match.group(1), match.group(2)))
    return answers


def _checkbox_answer(text: str, code: str) -> AnswerOption | None:
    match = re.match(r"^\s*[☐□☒✓✔]\s*(.+)$", normalize_space(text))
    if not match:
        return None
    return _make_answer(code, normalize_space(match.group(1)))


def _plain_paragraph_answers(
    window: list[Block], question_text: str, question_type: QuestionType
) -> list[AnswerOption]:
    """Read answer lists formatted as consecutive unnumbered paragraphs."""
    if question_type not in {QuestionType.SINGLE, QuestionType.MULTI, QuestionType.SCALE, QuestionType.UNKNOWN}:
        return []

    type_marker = re.compile(
        r"ОДИН\s+(?:ВАРИАНТ\s+)?ОТВЕТ|НЕСКОЛЬК\w*\s+(?:ВАРИАНТ\w*\s+)?ОТВЕТ|"
        r"МНОЖЕСТВЕН\w*\s+(?:ВЫБОР|ОТВЕТ)",
        flags=re.I,
    )
    start_index: int | None = 1 if type_marker.search(question_text) else None
    if start_index is None:
        for index, block in enumerate(window[1:], start=1):
            if block.kind == "paragraph" and type_marker.search(block.text):
                start_index = index + 1
                break
    if start_index is None:
        # Some supplied questionnaires omit the "one answer" note entirely.
        # For an otherwise unknown question, two or more consecutive response
        # lines are strong enough evidence of a positional single-choice list.
        if question_type == QuestionType.UNKNOWN:
            start_index = 1
        else:
            return []

    answers: list[AnswerOption] = []
    for block in window[start_index:]:
        if block.kind != "paragraph":
            if answers:
                break
            continue
        value = normalize_space(block.text)
        upper = value.upper().replace("Ё", "Е")
        if not value:
            continue
        if looks_like_instruction(value):
            if answers:
                break
            continue
        if re.match(r"^[A-Za-zА-Яа-яЁё]\.(?:\s|$)", value):
            break
        if re.match(r"^(?:PSM|БЛОК|СКРИНИНГ|МЕТОДОЛОГИЯ|ИНФОБЛОК)\b", upper):
            break
        if upper.startswith(("ПОЛЕ ДЛЯ", "ОТКРЫТЫЙ ОТВЕТ", "ПОКАЗАТЬ ИЗОБРАЖЕНИЕ", "ВЫВЕСТИ ИЗОБРАЖЕНИЕ")):
            break
        if re.fullmatch(
            r"[A-Za-zА-Яа-яЁё]\d+(?:[._]\d+)?(?:\s*[<>=]+\s*[A-Za-zА-Яа-яЁё]\d+(?:[._]\d+)?)+",
            value,
        ):
            break
        if len(value) <= 90 and value == upper and any(char.isalpha() for char in value):
            break
        answers.append(_make_answer(str(len(answers) + 1), value))
    if question_type == QuestionType.UNKNOWN and len(answers) < 2:
        return []
    return answers


def _positional_matrix_answers(
    window: list[Block], ignored_table_ids: set[int] | None = None
) -> tuple[list[AnswerOption], list[AnswerOption]]:
    """Read matrices whose Word mock-up uses checkbox glyphs but no codes."""
    groups: dict[int, list[Block]] = {}
    for block in window:
        if block.table_id is not None and block.table_id not in (ignored_table_ids or set()):
            groups.setdefault(block.table_id, []).append(block)
    for rows in groups.values():
        if len(rows) < 2 or len(rows[0].cells) < 2:
            continue
        header = rows[0].cells
        data_rows = rows[1:]
        checkbox_grid = any(any(cell in {"☐", "□", "☒", "○", "◯", "●"} for cell in row.cells[1:]) for row in data_rows)
        if not checkbox_grid:
            continue
        columns: list[AnswerOption] = []
        for index, value in enumerate(header[1:], start=1):
            value = normalize_space(value)
            if not value:
                continue
            coded = re.match(r"^(\d{1,3})\s*(.*)$", value)
            code = coded.group(1) if coded else str(index)
            label = normalize_space(coded.group(2)) if coded and coded.group(2) else value
            columns.append(_make_answer(code, label))
        matrix_rows = [
            _make_answer(str(index), normalize_space(row.cells[0]))
            for index, row in enumerate(data_rows, start=1)
            if row.cells and normalize_space(row.cells[0])
        ]
        if columns or matrix_rows:
            return columns, matrix_rows
    return [], []


def _exclusive_codes(instructions: list[str]) -> set[str]:
    codes: set[str] = set()
    for instruction in instructions:
        upper = instruction.upper().replace("Ё", "Е")
        if "ИСКЛЮЧА" not in upper:
            continue
        for match in re.finditer(r"(?:КОД(?:Ы)?\s*)?(\d{1,3})(?:\s*[-–—]\s*(\d{1,3}))?", upper):
            start = int(match.group(1))
            end = int(match.group(2) or start)
            if 0 <= end - start <= 20:
                codes.update(str(item) for item in range(start, end + 1))
    return codes


def _technical_table_ids(blocks: list[Block]) -> set[int]:
    groups: dict[int, list[tuple[int, Block]]] = {}
    for position, block in enumerate(blocks):
        if block.table_id is not None:
            groups.setdefault(block.table_id, []).append((position, block))
    ignored: set[int] = set()
    for table_id, entries in groups.items():
        first_position = entries[0][0]
        table_text = " ".join(item.raw_text or item.text for _, item in entries[:3]).upper().replace("Ё", "Е")
        previous_texts = [
            blocks[index].raw_text or blocks[index].text
            for index in range(max(0, first_position - 3), first_position)
            if blocks[index].table_id is None
        ]
        previous = " ".join(previous_texts).upper().replace("Ё", "Е")
        if (
            ("ТЕСТИРУЕМ" in table_text and "КОД В ПРОГРАММ" in table_text)
            or re.search(r"СУПЕРВАЙЗЕР.{0,30}(?:РОТАЦ|ЦЕН)", previous)
        ):
            ignored.add(table_id)
    return ignored


def _inline_labels_by_question(text: str) -> dict[str, str]:
    matches = list(
        re.finditer(rf"(?:^|\s)({QUESTION_ID_TOKEN})\s*:\s*", text, flags=re.I)
    )
    labels: dict[str, str] = {}
    for index, match in enumerate(matches):
        end = matches[index + 1].start() if index + 1 < len(matches) else len(text)
        label = normalize_space(text[match.end():end])
        if label:
            labels[normalize_id(match.group(1))] = label
    return labels


def _test_products(table_groups: dict[int, list[Block]]) -> list[dict[str, str]]:
    products: list[dict[str, str]] = []
    for rows in table_groups.values():
        if not rows:
            continue
        header = " ".join(rows[0].raw_cells or rows[0].cells).upper().replace("Ё", "Е")
        if "ТЕСТИРУЕМ" not in header or "КОД В ПРОГРАММ" not in header:
            continue
        header_cells = [normalize_space(cell).upper().replace("Ё", "Е") for cell in rows[0].cells]
        name_column = next((index for index, cell in enumerate(header_cells) if "ТЕСТИРУЕМ" in cell), 0)
        code_column = next((index for index, cell in enumerate(header_cells) if "КОД В ПРОГРАММ" in cell), 1)
        for row in rows[1:]:
            cells = row.cells
            if max(name_column, code_column) >= len(cells):
                continue
            name = normalize_space(cells[name_column])
            code = normalize_space(cells[code_column])
            if name and code:
                products.append({"name": name, "code": code})
    return products


class DocxSurveyParser:
    def parse(self, path: str | Path) -> SurveySpec:
        path = Path(path)
        blocks = extract_blocks(path)
        ignored_table_ids = _technical_table_ids(blocks)
        title = next((b.text for b in blocks if b.text), path.stem)
        question_hits: list[tuple[int, str, str]] = []
        question_boundaries = [position for position, block in enumerate(blocks) if _is_question_boundary(block)]
        for position, block in enumerate(blocks):
            parsed = _question_from_block(block)
            if parsed:
                question_hits.append((position, parsed[0], parsed[1]))

        questions: list[QuestionSpec] = []
        warnings: list[str] = []
        seen_ids: dict[str, int] = {}
        for hit_index, (position, qid, text) in enumerate(question_hits):
            next_position = next((item for item in question_boundaries if item > position), len(blocks))
            window = blocks[position:next_position]
            options: list[AnswerOption] = []
            matrix_rows: list[AnswerOption] = []
            instructions: list[str] = []
            rotation = None
            for block in window[1:]:
                answers = [] if block.table_id in ignored_table_ids else _answers_from_block(block)
                for answer in answers:
                    if not any(item.code == answer.code and item.text == answer.text for item in options):
                        options.append(answer)
                if looks_like_instruction(block.text):
                    instructions.append(block.text)
                candidate_rotation = parse_rotation(block.text)
                if candidate_rotation:
                    rotation = candidate_rotation

            exclusive_codes = _exclusive_codes(instructions)
            for option in options:
                if option.code in exclusive_codes:
                    option.exclusive = True

            local_context = " ".join(
                block.text for block in window[:8] if block.kind == "paragraph"
            )
            # First look for an explicit type in the wording/instructions;
            # only after that use the presence of options as a SINGLE fallback.
            inferred_type = _infer_type(text, False)
            if inferred_type == QuestionType.UNKNOWN:
                inferred_type = _infer_type(local_context, False)
            if inferred_type == QuestionType.UNKNOWN:
                inferred_type = _infer_type(local_context, bool(options))

            if not options:
                for block in window[1:]:
                    if block.kind != "paragraph":
                        continue
                    answer = _checkbox_answer(block.text, str(len(options) + 1))
                    if answer:
                        options.append(answer)

            if not options:
                options = _plain_paragraph_answers(window, text, inferred_type)
                if options and inferred_type == QuestionType.UNKNOWN:
                    inferred_type = QuestionType.SINGLE

            if not options and inferred_type == QuestionType.RANKING:
                for block in window[1:]:
                    if block.kind != "paragraph" or looks_like_instruction(block.text):
                        continue
                    value = normalize_space(block.text)
                    if value:
                        options.append(_make_answer(str(len(options) + 1), value))

            positional_options, positional_rows = _positional_matrix_answers(window, ignored_table_ids)
            if positional_rows:
                if inferred_type not in {QuestionType.MATRIX_SINGLE, QuestionType.MATRIX_MULTI}:
                    inferred_type = QuestionType.MATRIX_SINGLE
                existing_by_code = {option.code: option for option in options}
                ordered_options = [existing_by_code.get(option.code, option) for option in positional_options]
                header_codes = {option.code for option in positional_options}
                ordered_options.extend(option for option in options if option.code not in header_codes)
                options = ordered_options
                matrix_rows = positional_rows
            seen_ids[qid] = seen_ids.get(qid, 0) + 1
            if seen_ids[qid] > 1:
                warnings.append(f"Код вопроса {qid} встречается несколько раз; возможно, это цикл или блок по концептам.")
            questions.append(
                QuestionSpec(
                    id=qid,
                    text=text,
                    type=inferred_type,
                    options=options,
                    matrix_rows=matrix_rows,
                    rotation=rotation,
                    instructions=instructions,
                    source=SourceLocation(
                        block_index=blocks[position].index,
                        kind=blocks[position].kind,
                        style=blocks[position].style or None,
                    ),
                )
            )

        by_id: dict[str, list[QuestionSpec]] = {}
        for question in questions:
            by_id.setdefault(question.id, []).append(question)
        table_groups: dict[int, list[Block]] = {}
        for block in blocks:
            if block.table_id is not None:
                table_groups.setdefault(block.table_id, []).append(block)
        for rows in table_groups.values():
            if not rows:
                continue
            if rows[0].table_id in ignored_table_ids:
                continue
            header = rows[0].cells
            header_ids = [
                normalize_id(cell)
                if re.fullmatch(QUESTION_ID_TOKEN, normalize_space(cell))
                else ""
                for cell in header
            ]
            linked_columns = [(index, qid) for index, qid in enumerate(header_ids) if qid in by_id]
            if not linked_columns:
                continue
            linked_indexes = {index for index, _ in linked_columns}
            label_candidates = [index for index in range(len(header)) if index not in linked_indexes]
            label_column = max(
                label_candidates,
                key=lambda column: sum(
                    1
                    for row in rows[1:]
                    if column < len(row.cells)
                    and any(char.isalpha() for char in normalize_space(row.cells[column]))
                ),
                default=0,
            )
            for column, qid in linked_columns:
                target = by_id[qid][0]
                for row in rows[1:]:
                    if not row.cells or column >= len(row.cells):
                        continue
                    base_text = normalize_space(row.cells[label_column]) if label_column < len(row.cells) else ""
                    inline_labels = _inline_labels_by_question(base_text)
                    text = inline_labels.get(qid, base_text)
                    # Some shared Word tables put a generic last label in one
                    # row and the question-specific replacement labels in the
                    # following row (for example S8/S9). The second row is an
                    # override, not an additional answer or a new question.
                    if len(inline_labels) > 1 and qid in inline_labels:
                        if target.options:
                            target.options[-1].text = text
                        continue
                    code = normalize_space(row.cells[column])
                    if text and not ANSWER_CODE_PATTERN.match(code):
                        if qid in inline_labels:
                            fallback_codes = [
                                normalize_space(row.cells[index])
                                for index in sorted(linked_indexes)
                                if index < len(row.cells) and ANSWER_CODE_PATTERN.match(normalize_space(row.cells[index]))
                            ]
                            if fallback_codes:
                                code = fallback_codes[-1]
                    if text and ANSWER_CODE_PATTERN.match(code):
                        answer = _make_answer(code, text)
                        if not any(item.code == answer.code and item.text == answer.text for item in target.options):
                            target.options.append(answer)
                exclusive_codes = _exclusive_codes(target.instructions)
                for option in target.options:
                    if option.code in exclusive_codes:
                        option.exclusive = True

        question_positions = [(position, question) for (position, _, _), question in zip(question_hits, questions)]
        question_block_positions = set(question_boundaries)
        rules = []
        for position, block in enumerate(blocks):
            # Question wording often contains words like "если" or
            # "продолжите" and may end with "ротация строк".  It is still
            # question content, not a routing instruction.
            if position in question_block_positions:
                continue
            block_rules = parse_logic_rules(block.text, source_block=block.index)
            for rule in block_rules:
                # Quotas are deliberately outside the scope of this checker.
                # A sentence such as "иначе завершить. Поставить квоты..." may
                # leave a condition-less fragment after instruction splitting;
                # it must not become a survey-flow assertion.
                if rule.condition is None and "КВОТ" in rule.raw_text.upper():
                    continue

                # Formatting notes about exclusivity and similar metadata are
                # consumed while parsing answer rows, not as navigation rules.
                if rule.action == "other" and rule.condition is None:
                    continue

                if not rule.target_question_id and rule.action in {"show", "hide"}:
                    next_questions = [question for qpos, question in question_positions if qpos > position]
                    if next_questions:
                        rule.target_question_id = next_questions[0].id
                        rule.confidence = min(rule.confidence, 0.82)
                if rule.condition is None and "ЕСЛИ" in rule.raw_text.upper():
                    previous_questions = [question for qpos, question in question_positions if qpos < position]
                    code_match = re.search(
                        r"ЕСЛИ\s+(?P<negative>НЕ\s+)?(?:ОТМЕЧЕНЫ?|ВЫБРАНЫ?|ВЫБРАН\s+ОДИН\s+ИЗ)\s+"
                        r"КОД(?:А|Ы|ОВ)?\s+(?P<codes>[\d\s,;–—-]+)",
                        rule.raw_text,
                        flags=re.I,
                    )
                    explicit_codes = re.findall(r"КОД(?:А|Ы|ОВ)?\s*(\d+)", rule.raw_text, flags=re.I)
                    if previous_questions and (code_match or explicit_codes):
                        from .models import ConditionAtom, ConditionGroup
                        from .normalization import parse_code_list

                        values = parse_code_list(code_match.group("codes")) if code_match else explicit_codes
                        if values:
                            rule.condition = ConditionGroup(
                                combinator="and",
                                atoms=[
                                    ConditionAtom(
                                        question_id=previous_questions[-1].id,
                                        operator="not_in" if code_match and code_match.group("negative") else "in",
                                        values=values,
                                    )
                                ],
                            )
                            rule.confidence = 0.82
                if rule.target_question_id:
                    target = next((q for q in questions if q.id == rule.target_question_id), None)
                    if target and rule.action in {"show", "hide"}:
                        target.display_rules.append(rule)
                rules.append(rule)
                if rule.condition is None and "ЕСЛИ" in rule.raw_text.upper():
                    warnings.append(f"Не удалось однозначно разобрать условие: {rule.raw_text}")

        if not questions:
            warnings.append("В документе не найдено ни одного вопроса с кодом.")
        return SurveySpec(
            title=title,
            questions=questions,
            rules=rules,
            warnings=list(dict.fromkeys(warnings)),
            metadata={
                "source_filename": path.name,
                "block_count": len(blocks),
                "struck_or_deleted_fragments_ignored": True,
                "ignored_technical_table_ids": sorted(ignored_table_ids),
                "test_products": _test_products(table_groups),
                "piped_options": {
                    normalize_id(match.group("target")): normalize_id(match.group("source"))
                    for block in blocks
                    for match in [
                        re.search(
                            rf"ДЛЯ\s+(?P<target>{QUESTION_ID_TOKEN})\s+"
                            rf"ПОКАЗАТЬ.+?ОТМЕЧЕНН\w*\s+В\s+(?P<source>{QUESTION_ID_TOKEN})",
                            block.text,
                            flags=re.I,
                        )
                    ]
                    if match
                },
            },
        )
