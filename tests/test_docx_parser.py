import tempfile
import unittest
from pathlib import Path

from docx import Document

from app.docx_parser import DocxSurveyParser
from app.models import QuestionType


class DocxParserTests(unittest.TestCase):
    def _build_sample(self, path: Path) -> None:
        document = Document()
        document.add_heading("Тестовая анкета", level=1)
        document.add_paragraph("S1. Какие продукты вы покупали?")
        document.add_paragraph("Несколько ответов")
        table = document.add_table(rows=1, cols=2)
        table.rows[0].cells[0].text = "1"
        table.rows[0].cells[1].text = "Бургеры"
        row = table.add_row()
        row.cells[0].text = "99"
        row.cells[1].text = "Ничего из перечисленного"
        document.add_paragraph("Код 99 — исключающий")
        document.add_paragraph("Q1. Оцените идею")
        document.add_paragraph("Один ответ")
        table = document.add_table(rows=1, cols=2)
        table.rows[0].cells[0].text = "1"
        table.rows[0].cells[1].text = "Не нравится"
        row = table.add_row()
        row.cells[0].text = "5"
        row.cells[1].text = "Очень нравится"
        document.add_paragraph("Q2. Почему идея вам нравится?")
        document.add_paragraph("Открытый ответ")
        document.add_paragraph("Показать Q2, если в Q1 выбран код 5")
        document.save(path)

    def test_reads_questions_options_types_and_logic(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "sample.docx"
            self._build_sample(path)
            spec = DocxSurveyParser().parse(path)

        self.assertEqual([question.id for question in spec.questions], ["S1", "Q1", "Q2"])
        self.assertEqual(spec.questions[0].type, QuestionType.MULTI)
        self.assertTrue(spec.questions[0].options[-1].exclusive)
        self.assertEqual(spec.questions[2].type, QuestionType.TEXT)
        self.assertEqual(len(spec.questions[2].display_rules), 1)
        self.assertEqual(spec.questions[2].display_rules[0].condition.atoms[0].values, ["5"])

    def test_reads_uncoded_paragraph_options_and_ignores_formula(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "plain-options.docx"
            document = Document()
            document.add_paragraph("S1. Выберите ресторан")
            document.add_paragraph("Один ответ")
            document.add_paragraph("Бургер Кинг")
            document.add_paragraph("Другой ресторан")
            document.add_paragraph("E5. При какой цене блюдо покажется дорогим?")
            document.add_paragraph("E4 < E5 < E2")
            document.save(path)

            spec = DocxSurveyParser().parse(path)

        self.assertEqual([option.code for option in spec.questions[0].options], ["1", "2"])
        self.assertEqual(spec.questions[0].type, QuestionType.SINGLE)
        self.assertEqual(spec.questions[1].type, QuestionType.NUMBER)
        self.assertEqual(spec.questions[1].options, [])

    def test_infers_previous_question_for_code_only_screening_rule(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "screening.docx"
            document = Document()
            document.add_paragraph("S1. Выберите ресторан")
            document.add_paragraph("Один ответ")
            document.add_paragraph("Бургер Кинг")
            document.add_paragraph("Другой ресторан")
            document.add_paragraph("Закончить интервью, если не выбран код 1")
            document.save(path)

            spec = DocxSurveyParser().parse(path)

        rule = next(rule for rule in spec.rules if rule.action == "terminate")
        self.assertEqual(rule.condition.atoms[0].question_id, "S1")
        self.assertEqual(rule.condition.atoms[0].operator, "not_in")
        self.assertEqual(rule.condition.atoms[0].values, ["1"])

    def test_ignores_full_and_partial_strikethrough(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "strikethrough.docx"
            document = Document()
            question = document.add_paragraph()
            question.add_run("Q1. Какой ")
            removed = question.add_run("бургер ")
            removed.font.strike = True
            question.add_run("ролл Вам нравится?")
            document.add_paragraph("Один ответ")
            table = document.add_table(rows=1, cols=2)
            table.rows[0].cells[0].text = "1"
            table.rows[0].cells[1].text = "Чикен ролл"
            row = table.add_row()
            row.cells[0].text = "2"
            removed_option = row.cells[1].paragraphs[0].add_run("Воппер")
            removed_option.font.strike = True

            removed_rule = document.add_paragraph()
            rule_run = removed_rule.add_run("Показать Q3, если в Q1 выбран код 1")
            rule_run.font.strike = True

            removed_question = document.add_paragraph()
            removed_question_run = removed_question.add_run("Q2. Удалённый вопрос?")
            removed_question_run.font.strike = True
            document.add_paragraph("Открытый ответ")
            document.add_paragraph("Этот текст относился к удалённому вопросу")
            document.add_paragraph("Q3. Активный вопрос?")
            document.add_paragraph("Открытый ответ")
            document.save(path)

            spec = DocxSurveyParser().parse(path)

        self.assertEqual([item.id for item in spec.questions], ["Q1", "Q3"])
        self.assertEqual(spec.questions[0].text, "Какой ролл Вам нравится?")
        self.assertEqual([item.text for item in spec.questions[0].options], ["Чикен ролл"])
        self.assertEqual(spec.rules, [])
        self.assertNotIn("Этот текст относился", " ".join(spec.questions[0].instructions))

    def test_reads_shared_answer_tables_and_inline_last_labels(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "shared-tables.docx"
            document = Document()
            document.add_paragraph("S5. Какие марки Вы покупали?")
            document.add_paragraph("Несколько ответов")
            document.add_paragraph("Для S6 показать марки, отмеченные в S5")
            document.add_paragraph("S6. Какую марку Вы покупаете чаще всего?")
            document.add_paragraph("Один ответ")
            document.add_paragraph("S7. Какие марки Вы знаете?")
            brands = document.add_table(rows=1, cols=5)
            for index, value in enumerate(["", "", "S5", "S6", "S7"]):
                brands.rows[0].cells[index].text = value
            for code, brand in [("1", "Burger King"), ("2", "Ростикс")]:
                row = brands.add_row()
                for index, value in enumerate([code, brand, code, code, code]):
                    row.cells[index].text = value

            document.add_paragraph("S8. Что Вы покупали?")
            document.add_paragraph("S9. Что Вы не будете покупать?")
            products = document.add_table(rows=1, cols=3)
            for index, value in enumerate(["", "S8", "S9"]):
                products.rows[0].cells[index].text = value
            row = products.add_row()
            for index, value in enumerate(["Бургеры", "1", "1"]):
                row.cells[index].text = value
            row = products.add_row()
            for index, value in enumerate(["Ничего из перечисленного", "98", "98"]):
                row.cells[index].text = value
            row = products.add_row()
            for index, value in enumerate(
                [
                    "s8: Ничего из перечисленного не покупал(а) s9: Ничего из перечисленного не буду покупать",
                    "-",
                    "99",
                ]
            ):
                row.cells[index].text = value
            document.save(path)

            spec = DocxSurveyParser().parse(path)

        by_id = {question.id: question for question in spec.questions}
        self.assertEqual([option.text for option in by_id["S5"].options], ["Burger King", "Ростикс"])
        self.assertEqual([option.text for option in by_id["S6"].options], ["Burger King", "Ростикс"])
        self.assertEqual([option.text for option in by_id["S7"].options], ["Burger King", "Ростикс"])
        self.assertEqual(
            [(option.code, option.text) for option in by_id["S8"].options],
            [("1", "Бургеры"), ("98", "Ничего из перечисленного не покупал(а)")],
        )
        self.assertEqual(
            [(option.code, option.text) for option in by_id["S9"].options],
            [("1", "Бургеры"), ("98", "Ничего из перечисленного не буду покупать")],
        )
        self.assertEqual(spec.metadata["piped_options"], {"S6": "S5"})
        self.assertEqual([question.id for question in spec.questions].count("S8"), 1)

    def test_ignores_rotation_and_product_mapping_tables(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "technical-tables.docx"
            document = Document()
            document.add_paragraph("S11. Вы согласны попробовать напитки?")
            document.add_paragraph("Один ответ")
            answers = document.add_table(rows=1, cols=2)
            answers.rows[0].cells[0].text = "1"
            answers.rows[0].cells[1].text = "Да"
            row = answers.add_row()
            row.cells[0].text = "2"
            row.cells[1].text = "Нет"
            document.add_paragraph("СУПЕРВАЙЗЕР, ОТМЕТЬТЕ РОТАЦИЮ")
            rotation = document.add_table(rows=1, cols=2)
            rotation.rows[0].cells[0].text = "1"
            rotation.rows[0].cells[1].text = "195 – 349"
            row = rotation.add_row()
            row.cells[0].text = "2"
            row.cells[1].text = "349 – 195"

            document.add_paragraph("P10. Как часто Вы покупаете напиток?")
            document.add_paragraph("Один ответ")
            frequency = document.add_table(rows=1, cols=2)
            frequency.rows[0].cells[0].text = "1"
            frequency.rows[0].cells[1].text = "Раз в неделю"
            row = frequency.add_row()
            row.cells[0].text = "2"
            row.cells[1].text = "Реже"
            products = document.add_table(rows=1, cols=3)
            for index, value in enumerate(["F4", "Тестируемые продукты", "Код в программе"]):
                products.rows[0].cells[index].text = value
            row = products.add_row()
            for index, value in enumerate(["1", "Шоколадный милкшейк БК", "Напиток 195"]):
                row.cells[index].text = value
            document.save(path)

            spec = DocxSurveyParser().parse(path)

        by_id = {question.id: question for question in spec.questions}
        self.assertEqual([option.text for option in by_id["S11"].options], ["Да", "Нет"])
        self.assertEqual([option.text for option in by_id["P10"].options], ["Раз в неделю", "Реже"])
        self.assertEqual(
            spec.metadata["test_products"],
            [{"name": "Шоколадный милкшейк БК", "code": "Напиток 195"}],
        )

    def test_explicit_one_answer_wins_over_option_wording(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "one-answer.docx"
            document = Document()
            document.add_paragraph("A1_2. Как часто Вы покупаете напитки?")
            document.add_paragraph("ОДИН ОТВЕТ")
            table = document.add_table(rows=1, cols=2)
            table.rows[0].cells[0].text = "1"
            table.rows[0].cells[1].text = "Несколько раз в неделю"
            row = table.add_row()
            row.cells[0].text = "2"
            row.cells[1].text = "Несколько раз в месяц"
            document.save(path)

            spec = DocxSurveyParser().parse(path)

        self.assertEqual(spec.questions[0].type, QuestionType.SINGLE)

    def test_reads_explicit_info_and_intro_text_as_questions(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "info-blocks.docx"
            document = Document()
            document.add_paragraph("INFO: Внимательно посмотрите на описание бургера.")
            document.add_paragraph(
                "Вводный текст для концепции: Далее мы зададим несколько вопросов о Вашем впечатлении."
            )
            document.add_paragraph("Q1. Понравилась ли Вам концепция?")
            document.add_paragraph("Один ответ")
            document.save(path)

            spec = DocxSurveyParser().parse(path)

        self.assertEqual([question.id for question in spec.questions], ["INFO1", "INFO2", "Q1"])
        self.assertEqual(spec.questions[0].type, QuestionType.INFO)
        self.assertEqual(spec.questions[0].text, "Внимательно посмотрите на описание бургера.")
        self.assertEqual(
            spec.questions[1].text,
            "Далее мы зададим несколько вопросов о Вашем впечатлении.",
        )

    def test_reads_monadic_catalogue_scoped_questions_and_dynamic_wording(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "concepts.docx"
            document = Document()
            document.add_paragraph("Каждый респондент случайно видит только одну концепцию")
            document.add_paragraph("ЕСЛИ тестируется Прайм Чикен на сырной булочке")
            document.add_paragraph("Q6.1. Нравится ли Вам булочка?")
            document.add_paragraph("ЕСЛИ тестируется Прайм Чикен на вафле (v2-3)")
            document.add_paragraph("Q6.2. Нравится ли Вам вафля?")
            document.add_paragraph("ЕСЛИ Q6.2=1-3")
            document.add_paragraph("Q7. Почему не понравилась вафля?")
            document.add_paragraph("Формулировка в зависимости от концепта")
            document.add_paragraph("Q8. Оцените основу (булочке или вафле)")
            document.save(path)

            spec = DocxSurveyParser().parse(path)

        self.assertEqual(len(spec.metadata["concepts"]), 3)
        self.assertEqual(spec.metadata["direct_concept_question_ids"]["Q6.1"], ["concept-1"])
        self.assertEqual(
            spec.metadata["direct_concept_question_ids"]["Q6.2"],
            ["concept-2", "concept-3"],
        )
        by_id = {question.id: question for question in spec.questions}
        self.assertEqual(by_id["Q7"].concept_ids, ["concept-2", "concept-3"])
        self.assertEqual(spec.metadata["dynamic_concept_question_ids"], ["Q8"])

    def test_reads_unlabelled_intro_before_concept_image_as_info(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "implicit-info.docx"
            document = Document()
            document.add_paragraph(
                "Представьте, что в меню появилась новая линейка блюд. Сейчас Вы увидите "
                "ее подробное описание, после чего мы попросим Вас ответить на вопросы."
            )
            document.add_paragraph("Показать изображение и описание")
            document.add_paragraph("Q1. Нравится ли Вам линейка?")
            document.save(path)

            spec = DocxSurveyParser().parse(path)

        self.assertEqual(spec.questions[0].id, "INFO1")
        self.assertEqual(spec.questions[0].type, QuestionType.INFO)

    def test_reads_numbered_info_as_info_not_multi(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "numbered-info.docx"
            document = Document()
            document.add_paragraph(
                "INFO1. Сейчас мы покажем несколько названий и зададим несколько вопросов."
            )
            document.add_paragraph("Q1. Какое название Вам нравится?")
            document.add_paragraph("Один ответ")
            document.save(path)

            spec = DocxSurveyParser().parse(path)

        self.assertEqual(spec.questions[0].id, "INFO1")
        self.assertEqual(spec.questions[0].type, QuestionType.INFO)

    def test_only_explicit_exclusive_code_is_marked(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "exclusive-code.docx"
            document = Document()
            document.add_paragraph("S1. Какие места Вы посещали?")
            document.add_paragraph(
                "Тип ответа: несколько ответов, ротация кодов 1-7. "
                "Код 99 «Ничего» — исключающий."
            )
            table = document.add_table(rows=1, cols=2)
            table.rows[0].cells[0].text = "Код"
            table.rows[0].cells[1].text = "Вариант ответа"
            for code, text in [("1", "Ресторан"), ("7", "Кафе"), ("99", "Ничего")]:
                row = table.add_row()
                row.cells[0].text = code
                row.cells[1].text = text
            document.save(path)

            spec = DocxSurveyParser().parse(path)

        self.assertEqual(
            [(option.code, option.exclusive) for option in spec.questions[0].options],
            [("1", False), ("7", False), ("99", True)],
        )

    def test_post_question_show_only_instruction_targets_previous_question(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "post-question-filter.docx"
            document = Document()
            document.add_paragraph("B5. Выберите название")
            document.add_paragraph("Один ответ")
            document.add_paragraph("B6. Что бы Вы сделали?")
            document.add_paragraph("Несколько ответов")
            document.add_paragraph("Показывать только если B5 НЕ равно коду 99")
            document.add_paragraph("A2. Следующий безусловный вопрос")
            document.save(path)

            spec = DocxSurveyParser().parse(path)

        by_id = {question.id: question for question in spec.questions}
        self.assertEqual(len(by_id["B6"].display_rules), 1)
        self.assertEqual(by_id["A2"].display_rules, [])
        atom = by_id["B6"].display_rules[0].condition.atoms[0]
        self.assertEqual((atom.question_id, atom.operator, atom.values), ("B5", "not_in", ["99"]))

    def test_reads_separate_matrix_row_and_column_tables(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "sectioned-matrix.docx"
            document = Document()
            document.add_paragraph("S5. Насколько Вам подходят эти фразы?")
            document.add_paragraph(
                "Тип ответа: матрица. Строки ротировать. Один ответ в каждой строке."
            )
            document.add_paragraph("Строки матрицы")
            rows = document.add_table(rows=1, cols=2)
            rows.rows[0].cells[0].text = "Код"
            rows.rows[0].cells[1].text = "Строка"
            for code, text in [("1", "Первая фраза"), ("2", "Вторая фраза")]:
                row = rows.add_row()
                row.cells[0].text = code
                row.cells[1].text = text
            document.add_paragraph("Столбцы матрицы")
            columns = document.add_table(rows=1, cols=2)
            columns.rows[0].cells[0].text = "Код"
            columns.rows[0].cells[1].text = "Столбец / шкала"
            for code, text in [("1", "Совсем не про меня"), ("5", "Полностью про меня")]:
                row = columns.add_row()
                row.cells[0].text = code
                row.cells[1].text = text

            document.add_paragraph("B1. Какие названия подходят характеристикам?")
            document.add_paragraph(
                "Тип ответа: матрица с множественным выбором. Строки и столбцы ротировать, "
                "кроме кода 99. Код 99 исключающий."
            )
            document.add_paragraph("Строки")
            rows = document.add_table(rows=1, cols=2)
            rows.rows[0].cells[0].text = "Код"
            rows.rows[0].cells[1].text = "Строка"
            for code, text in [("1", "Необычное"), ("2", "Скучное")]:
                row = rows.add_row()
                row.cells[0].text = code
                row.cells[1].text = text
            document.add_paragraph("Столбцы")
            columns = document.add_table(rows=1, cols=2)
            columns.rows[0].cells[0].text = "Код"
            columns.rows[0].cells[1].text = "Столбец / шкала"
            for code, text in [("1", "Мохито"), ("2", "Мята-Лайм"), ("99", "Ни одно")]:
                row = columns.add_row()
                row.cells[0].text = code
                row.cells[1].text = text
            document.save(path)

            spec = DocxSurveyParser().parse(path)

        by_id = {question.id: question for question in spec.questions}
        self.assertEqual(by_id["S5"].type, QuestionType.MATRIX_SINGLE)
        self.assertEqual([option.text for option in by_id["S5"].options], ["Совсем не про меня", "Полностью про меня"])
        self.assertEqual([row.text for row in by_id["S5"].matrix_rows], ["Первая фраза", "Вторая фраза"])
        self.assertEqual(by_id["B1"].type, QuestionType.MATRIX_MULTI)
        self.assertEqual([option.text for option in by_id["B1"].options], ["Мохито", "Мята-Лайм", "Ни одно"])
        self.assertEqual([row.text for row in by_id["B1"].matrix_rows], ["Необычное", "Скучное"])
        self.assertEqual([option.exclusive for option in by_id["B1"].options], [False, False, True])


if __name__ == "__main__":
    unittest.main()
